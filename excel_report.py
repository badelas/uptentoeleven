from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


def flatten_dict(
    data,
    prefix=""
):
    rows = []

    if isinstance(data, dict):

        for key, value in data.items():

            new_key = (
                f"{prefix}.{key}"
                if prefix
                else key
            )

            rows.extend(
                flatten_dict(
                    value,
                    new_key
                )
            )

    elif isinstance(data, list):

        if all(
            not isinstance(
                value,
                (dict, list)
            )
            for value in data
        ):

            rows.append(
                (
                    prefix,
                    ", ".join(
                        str(value)
                        for value in data
                    )
                )
            )

        else:

            for index, value in enumerate(data):

                rows.extend(
                    flatten_dict(
                        value,
                        f"{prefix}[{index}]"
                    )
                )

    else:

        rows.append(
            (
                prefix,
                data
            )
        )

    return rows


def generate_excel_report(
    baseline: dict,
    findings: list,
    output_path: Path,
    comparison: list | None = None,
    readiness: list | None = None,
    target_family: str | None = None,
    target_version: str | None = None
):

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook = Workbook()

    # ========================================================
    # SUMMARY
    # ========================================================

    ws = workbook.active
    ws.title = "Summary"

    ws["A1"] = "Jira Upgrade Baseline Toolkit"
    ws["A1"].font = Font(
        bold=True,
        size=16
    )

    jira = baseline.get(
        "jira",
        {}
    )

    java = baseline.get(
        "java",
        {}
    )

    database = baseline.get(
        "database",
        {}
    )

    architecture = baseline.get(
        "architecture",
        {}
    )

    source_version = jira.get(
        "version"
    )

    target_display = (
        target_version
        if target_version
        else (
            f"{target_family}.x LTS"
            if target_family
            else "Non définie"
        )
    )

    architecture_display = (
        "Cluster"
        if architecture.get("clustered") is True
        else "Single node"
        if architecture.get("clustered") is False
        else "Non déterminée"
    )

    summary = [
        (
            "Instance",
            baseline.get(
                "instanceName",
                ""
            )
        ),
        (
            "Phase",
            baseline.get(
                "collectionPhase",
                ""
            )
        ),
        (
            "Jira Source",
            source_version
        ),
        (
            "Jira Cible",
            target_display
        ),
        (
            "Base URL",
            jira.get("baseUrl")
        ),
        (
            "Architecture",
            architecture_display
        ),
        (
            "Java Source",
            java.get("version")
        ),
        (
            "Database",
            database.get("type")
        ),
        (
            "Applications tierces",
            len(
                baseline.get(
                    "thirdPartyApps",
                    []
                )
            )
        )
    ]

    row = 3

    for label, value in summary:

        ws.cell(
            row=row,
            column=1,
            value=label
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=row,
            column=2,
            value=value
        )

        row += 1

    # ========================================================
    # BASELINE
    # ========================================================

    ws_baseline = workbook.create_sheet(
        "Baseline"
    )

    ws_baseline.append(
        [
            "Property",
            "Value"
        ]
    )

    for cell in ws_baseline[1]:
        cell.font = Font(
            bold=True
        )

    for key, value in flatten_dict(
        baseline
    ):

        ws_baseline.append(
            [
                key,
                (
                    str(value)
                    if value is not None
                    else ""
                )
            ]
        )

    # ========================================================
    # FINDINGS
    #
    # Etat de la SOURCE uniquement.
    # ========================================================

    ws_findings = workbook.create_sheet(
        "Findings"
    )

    ws_findings.append(
        [
            "Severity",
            "Domain",
            "Message"
        ]
    )

    for cell in ws_findings[1]:
        cell.font = Font(
            bold=True
        )

    for item in findings:

        ws_findings.append(
            [
                item.get("severity"),
                item.get("domain"),
                item.get("message")
            ]
        )

    # ========================================================
    # READINESS
    #
    # Analyse SOURCE -> CIBLE
    # ========================================================

    if readiness is not None:

        ws_readiness = workbook.create_sheet(
            "Readiness"
        )

        ws_readiness.append(
            [
                "Status",
                "Domain",
                "Check",
                "Current Value",
                "Target Value",
                "Message"
            ]
        )

        for cell in ws_readiness[1]:
            cell.font = Font(
                bold=True
            )

        for item in readiness:

            ws_readiness.append(
                [
                    item.get("status"),
                    item.get("domain"),
                    item.get("check"),
                    (
                        str(
                            item.get(
                                "current_value"
                            )
                        )
                        if item.get(
                            "current_value"
                        ) is not None
                        else ""
                    ),
                    (
                        str(
                            item.get(
                                "target_value"
                            )
                        )
                        if item.get(
                            "target_value"
                        ) is not None
                        else ""
                    ),
                    item.get("message")
                ]
            )

    # ========================================================
    # PRE / POST
    # ========================================================

    if comparison is not None:

        ws_compare = workbook.create_sheet(
            "PRE_POST"
        )

        ws_compare.append(
            [
                "Domain",
                "Field",
                "PRE",
                "POST",
                "Status"
            ]
        )

        for cell in ws_compare[1]:
            cell.font = Font(
                bold=True
            )

        for item in comparison:

            ws_compare.append(
                [
                    item.get("domain"),
                    item.get("field"),
                    (
                        str(item.get("pre"))
                        if item.get("pre") is not None
                        else ""
                    ),
                    (
                        str(item.get("post"))
                        if item.get("post") is not None
                        else ""
                    ),
                    item.get("status")
                ]
            )

    # ========================================================
    # LARGEUR DES COLONNES
    # ========================================================

    for sheet in workbook.worksheets:

        for column in sheet.columns:

            max_length = 0

            column_letter = (
                column[0].column_letter
            )

            for cell in column:

                value = cell.value

                if value is not None:

                    max_length = max(
                        max_length,
                        len(str(value))
                    )

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                80
            )

    workbook.save(
        output_path
    )

    return output_path